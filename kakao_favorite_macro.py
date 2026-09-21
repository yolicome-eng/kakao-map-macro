import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from pathlib import Path
import os, json, time, threading, subprocess, shutil, socket, struct, base64, re, html, urllib.request, urllib.parse
from openpyxl import load_workbook, Workbook

APP = Path(os.environ.get("APPDATA", str(Path.home()))) / "KakaoFavoriteMacro"
PROFILE = APP / "chrome_profile"
RESULT = APP / "result"
ERROR = APP / "error-excel"
STATE = APP / "state.json"
PORT = 9222
MAP = "https://map.kakao.com/"
for d in (APP, PROFILE, RESULT, ERROR): d.mkdir(parents=True, exist_ok=True)

def chrome():
    ps = [
        os.path.join(os.environ.get("PROGRAMFILES",""), "Google","Chrome","Application","chrome.exe"),
        os.path.join(os.environ.get("PROGRAMFILES(X86)",""), "Google","Chrome","Application","chrome.exe"),
        os.path.join(os.environ.get("LOCALAPPDATA",""), "Google","Chrome","Application","chrome.exe")
    ]
    return next((p for p in ps if os.path.exists(p)), shutil.which("chrome.exe") or shutil.which("chrome"))

def get_json(path):
    with urllib.request.urlopen("http://127.0.0.1:{}{}".format(PORT,path),timeout=5) as r:
        return json.loads(r.read().decode("utf-8"))

class WS:
    def __init__(self,url):
        u=urllib.parse.urlparse(url)
        self.s=socket.create_connection((u.hostname,u.port),timeout=10)
        k=base64.b64encode(os.urandom(16)).decode()
        req="GET {} HTTP/1.1\r\nHost: {}:{}\r\nUpgrade: websocket\r\nConnection: Upgrade\r\nSec-WebSocket-Key: {}\r\nSec-WebSocket-Version: 13\r\n\r\n".format(u.path or "/",u.hostname,u.port,k)
        self.s.sendall(req.encode())
        b=b""
        while b"\r\n\r\n" not in b: b+=self.s.recv(4096)
    def read(self,n):
        b=b""
        while len(b)<n:
            x=self.s.recv(n-len(b))
            if not x: raise RuntimeError("Chrome 연결이 종료되었습니다.")
            b+=x
        return b
    def send(self,o):
        raw=json.dumps(o,separators=(",",":")).encode(); L=len(raw)
        h=bytearray([129])
        if L<126: h.append(128|L)
        elif L<65536: h.extend([254]); h.extend(struct.pack("!H",L))
        else: h.extend([255]); h.extend(struct.pack("!Q",L))
        m=os.urandom(4); h.extend(m); h.extend(bytes(x^m[i%4] for i,x in enumerate(raw))); self.s.sendall(h)
    def recv(self):
        h=self.read(2); b2=h[1]; L=b2&127
        if L==126: L=struct.unpack("!H",self.read(2))[0]
        elif L==127: L=struct.unpack("!Q",self.read(8))[0]
        m=self.read(4) if b2&128 else b""; b=self.read(L)
        if m: b=bytes(x^m[i%4] for i,x in enumerate(b))
        if h[0]&15==1: return json.loads(b.decode())
        return self.recv()

class CDP:
    def __init__(self,url): self.w=WS(url); self.i=0
    def call(self,method,params=None,timeout=20):
        self.i+=1; i=self.i; self.w.send({"id":i,"method":method,"params":params or {}})
        end=time.time()+timeout
        while time.time()<end:
            r=self.w.recv()
            if r.get("id")==i:
                if "error" in r: raise RuntimeError(str(r["error"]))
                return r.get("result",{})
        raise TimeoutError(method)
    def js(self,x,timeout=25):
        r=self.call("Runtime.evaluate",{"expression":x,"awaitPromise":True,"returnByValue":True},timeout)
        return r.get("result",{}).get("value")
    def nav(self,url):
        self.call("Page.navigate",{"url":url},timeout=15); time.sleep(1)
    def close(self):
        try: self.w.s.close()
        except Exception: pass

def connect():
    for t in get_json("/json/list"):
        if t.get("type")=="page" and t.get("webSocketDebuggerUrl"):
            return CDP(t["webSocketDebuggerUrl"])

def start():
    p=chrome()
    if not p: raise RuntimeError("Google Chrome을 찾지 못했습니다.")
    subprocess.Popen([p,"--remote-debugging-port={}".format(PORT),"--remote-debugging-address=127.0.0.1","--remote-allow-origins=*","--user-data-dir={}".format(PROFILE),"--no-first-run","--no-default-browser-check",MAP],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
    for _ in range(40):
        time.sleep(.5)
        try:
            c=connect()
            if c: return c
        except Exception: pass
    raise RuntimeError("Chrome 연결에 실패했습니다. Chrome을 모두 종료한 후 다시 시도하세요.")


def js_quote(value):
    return json.dumps(str(value), ensure_ascii=False)

# Visible KakaoMap UI automation. This avoids the unstable mobile/internal
# search and coordinate endpoints used by the previous FIX3 build.
SEARCH_AND_SELECT_JS = r'''(async function(q){
 const sleep=ms=>new Promise(r=>setTimeout(r,ms));
 const clean=s=>String(s||'').replace(/\s+/g,' ').trim();
 const norm=s=>clean(s).replace(/[^0-9A-Za-z가-힣]/g,'').toLowerCase();
 const input=document.querySelector('#search\\.keyword\\.query');
 if(!input) return {ok:false,error:'카카오맵 검색창을 찾지 못했습니다.'};
 input.focus(); input.value=q;
 input.dispatchEvent(new Event('input',{bubbles:true}));
 input.dispatchEvent(new Event('change',{bubbles:true}));
 const btn=document.querySelector('#search\\.keyword\\.submit');
 if(btn) btn.click();
 else input.dispatchEvent(new KeyboardEvent('keydown',{key:'Enter',code:'Enter',keyCode:13,which:13,bubbles:true}));
 const deadline=Date.now()+6000;
 let items=[];
 while(Date.now()<deadline){
   items=[...document.querySelectorAll('.AddressItem')];
   if(!items.length) items=[...document.querySelectorAll('#info\\.search\\.place\\.list .AddressItem')];
   if(items.length) break;
   await sleep(150);
 }
 const nq=norm(q); let best=null,bestScore=-1;
 items.forEach((el,i)=>{
   const txt=clean(el.innerText||''), n=norm(txt); let score=0;
   if(n===nq) score+=1000;
   if(nq&&n.includes(nq)) score+=500;
   if(nq&&nq.includes(n)) score+=250;
   if(el.querySelector('.txt_address,.address')) score+=20;
   if(score>bestScore){bestScore=score;best={el,i,txt};}
 });
 if(!best) best={el:items[0],i:0,txt:items[0].innerText||''};
 best.el.scrollIntoView({block:'center'}); best.el.click(); await sleep(900);
 return {ok:true,text:clean(best.txt),index:best.i};
})(__Q__)'''

CLICK_FAVORITE_JS = r'''(async function(){
 const sleep=ms=>new Promise(r=>setTimeout(r,ms));
 for(let i=0;i<15;i++){
   const toolbar=document.querySelector('.InfoWindowToolbar');
   const fav=toolbar?toolbar.querySelector('.fav'):document.querySelector('.InfoWindowToolbar .fav');
   if(fav){fav.scrollIntoView({block:'center'});fav.click();await sleep(250);return {ok:true};}
   await sleep(150);
 }
 return {ok:false,error:'검색한 주소의 즐겨찾기 버튼을 찾지 못했습니다.'};
})()'''

SELECT_GROUP_JS = r'''(async function(group){
 const sleep=ms=>new Promise(r=>setTimeout(r,ms));
 const clean=s=>String(s||'').replace(/\s+/g,' ').trim();
 for(let i=0;i<15;i++){
   const body=clean(document.body.innerText||'');
   if(body.includes('이미 등록된 주소')) return {ok:true,duplicate:true};
   const strongs=[...document.querySelectorAll('strong.txt_folder')];
   const target=strongs.find(el=>clean(el.innerText)===clean(group));
   if(target){const a=target.closest('a')||target.parentElement;if(a){a.scrollIntoView({block:'center'});a.click();await sleep(250);return {ok:true};}}
   const folders=[...document.querySelectorAll('.list_folder')];
   const textTarget=folders.find(el=>clean(el.innerText).includes(clean(group)));
   if(textTarget){textTarget.click();await sleep(250);return {ok:true};}
   await sleep(150);
 }
 return {ok:false,error:'선택한 즐겨찾기 그룹을 찾지 못했습니다.'};
})(__GROUP__)'''

SAVE_FAVORITE_JS = r'''(async function(name){
 const sleep=ms=>new Promise(r=>setTimeout(r,ms));
 for(let i=0;i<15;i++){
   const input=document.querySelector('#display1');
   if(input){
     const setter=Object.getOwnPropertyDescriptor(HTMLInputElement.prototype,'value')?.set;
     input.focus(); if(setter) setter.call(input,name); else input.value=name;
     input.dispatchEvent(new Event('input',{bubbles:true}));
     input.dispatchEvent(new Event('change',{bubbles:true}));
     const ok=[...document.querySelectorAll('button')].find(b=>b.getAttribute('data-id')==='addOK'&&b.classList.contains('btn_submit'));
     if(ok){ok.click();await sleep(250);return {ok:true};}
   }
   await sleep(150);
 }
 return {ok:false,error:'즐겨찾기 저장 창을 찾지 못했습니다.'};
})(__NAME__)'''

FOLDERS = r'''fetch("/folder/list.json?sort=CREATE_AT").then(r=>r.json()).then(x=>x.result||x.folders||[]).catch(e=>[])'''

CLOSE_LAYERS_JS = r'''(()=>{[...document.querySelectorAll('.dimmedLayer')].forEach(e=>{try{e.remove()}catch(_){}});return true})()'''

class App:
    def __init__(self,root):
        self.r=root; self.c=None; self.rows=[]; self.folders=[]; self.stopflag=False; self.running=False
        root.title("카카오맵 주소 즐겨찾기 자동 등록"); root.geometry("780x620")
        self.ui(); self.load()
    def ui(self):
        f=ttk.Frame(self.r,padding=16); f.pack(fill="both",expand=True)
        ttk.Label(f,text="카카오맵 주소 즐겨찾기 자동 등록",font=("Malgun Gothic",18,"bold")).pack(anchor="w")
        ttk.Label(f,text="Chrome 로그인 후 엑셀 주소를 검색하여 선택한 그룹에 등록합니다.").pack(anchor="w",pady=5)
        bar=ttk.Frame(f); bar.pack(fill="x")
        ttk.Button(bar,text="① Chrome 시작 / 로그인",command=self.connect).pack(side="left")
        ttk.Button(bar,text="② 엑셀 선택",command=self.pick).pack(side="left",padx=8)
        self.file=ttk.Label(bar,text="파일: 선택 안 됨"); self.file.pack(side="left",fill="x",expand=True)
        g=ttk.LabelFrame(f,text="즐겨찾기 그룹",padding=10); g.pack(fill="x",pady=10)
        self.combo=ttk.Combobox(g,state="readonly",width=55); self.combo.pack(side="left")
        ttk.Button(g,text="그룹 새로고침",command=self.groups).pack(side="left",padx=8)
        o=ttk.LabelFrame(f,text="즐겨찾기 이름",padding=10); o.pack(fill="x")
        self.mode=tk.StringVar(value="space")
        ttk.Radiobutton(o,text="1~4열을 공백으로 연결",variable=self.mode,value="space").pack(anchor="w")
        ttk.Radiobutton(o,text="1~4열을 줄바꿈으로 연결",variable=self.mode,value="line").pack(anchor="w")
        self.pb=ttk.Progressbar(f,maximum=100); self.pb.pack(fill="x",pady=10)
        self.status=tk.StringVar(value="준비됨"); ttk.Label(f,textvariable=self.status).pack(anchor="w")
        self.log=tk.Text(f,height=15,font=("Consolas",9)); self.log.pack(fill="both",expand=True,pady=5)
        b=ttk.Frame(f); b.pack(fill="x")
        ttk.Button(b,text="▶ 등록 시작",command=self.begin).pack(side="left")
        ttk.Button(b,text="■ 중지",command=self.stop).pack(side="left",padx=8)
    def msg(self,s):
        self.r.after(0,lambda:(self.log.insert("end",s+"\n"),self.log.see("end")))
    def connect(self):
        def w():
            try:
                self.c=self.c or start(); self.msg("Chrome 연결 완료. Chrome에서 카카오 로그인/2차 인증을 완료하세요."); self.status.set("Chrome 연결됨")
            except Exception as e: self.r.after(0,lambda:messagebox.showerror("오류",str(e)))
        threading.Thread(target=w,daemon=True).start()
    def pick(self):
        p=filedialog.askopenfilename(filetypes=[("Excel","*.xlsx")])
        if not p:return
        try:
            ws=load_workbook(p,data_only=True).active
            if ws.max_column<5: raise ValueError("엑셀은 최소 5개 열이 필요합니다.")
            self.rows=[]
            for n,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
                if not any(str(x or "").strip() for x in row[:5]): continue
                self.rows.append([str(row[i] or "").strip() for i in range(5)])
            self.file.config(text="파일: {} / {}건".format(Path(p).name,len(self.rows))); self.msg("엑셀 {}건 불러옴".format(len(self.rows)))
        except Exception as e: messagebox.showerror("엑셀 오류",str(e))
    def groups(self):
        if not self.c: messagebox.showwarning("안내","먼저 Chrome을 시작하세요."); return
        def w():
            try:
                self.c.nav(MAP); self.folders=self.c.js(FOLDERS) or []
                names=["{} [ID {}]".format(x.get("title",""),x.get("folderId","")) for x in self.folders]
                self.r.after(0,lambda:self.combo.config(values=names))
                if names:self.r.after(0,lambda:self.combo.current(0))
                self.msg("그룹 {}개 확인".format(len(names)))
            except Exception as e:self.msg("그룹 조회 실패: "+str(e))
        threading.Thread(target=w,daemon=True).start()
    def load(self):
        try:
            self.last=json.loads(STATE.read_text(encoding="utf-8"))
        except Exception:self.last={}
    def stop(self): self.stopflag=True; self.status.set("중지 요청됨")
    def begin(self):
        if self.running:return
        if not self.c or not self.rows or not self.folders or self.combo.current()<0:
            messagebox.showwarning("안내","Chrome 시작 → 로그인 → 엑셀 선택 → 그룹 새로고침 순서로 진행하세요."); return
        if len(self.rows)>500 and not messagebox.askyesno("500건 초과","현재 {}건입니다. 한 그룹의 제한을 넘을 수 있습니다. 계속할까요?".format(len(self.rows))): return
        fid=str(self.folders[self.combo.current()].get("folderId"))
        threading.Thread(target=self.run,args=(fid,),daemon=True).start()
    def run(self,fid):
        self.running=True; self.stopflag=False
        total=len(self.rows); ok=dup=fail=0
        out=Workbook(); ow=out.active
        ow.append(['순번','선로명','선로번호','전산화번호','주소','상태','비고'])
        er=Workbook(); ew=er.active
        ew.append(['순번','선로명','선로번호','전산화번호','주소','오류'])
        group_name=self.combo.get().split(' [ID ')[0].strip()

        # 작업 시작 시 지도 페이지를 한 번만 준비합니다. 매 행마다 새로고침하지 않습니다.
        try:
            # 이전 실행에서 끊어진 CDP 소켓을 재사용하지 않고 새 연결을 확보합니다.
            fresh=connect()
            if fresh: self.c=fresh
            else: self.c=start()
            self.c.nav(MAP)
            time.sleep(2.0)
            ready_js=r'''(()=>({
                url:location.href,
                inputs:[...document.querySelectorAll('input')].map(x=>({id:x.id,cls:x.className,ph:x.placeholder,aria:x.getAttribute('aria-label')})).slice(0,30)
            }))()'''
            info=self.c.js(ready_js,10) or {}
            ready=bool(self.c.js(r'''!!(
                document.querySelector('#search\\.keyword\\.query') ||
                document.querySelector('input[name="q"]') ||
                document.querySelector('input[placeholder*="검색"]') ||
                document.querySelector('input[aria-label*="검색"]') ||
                [...document.querySelectorAll('input')].find(x=>/search|keyword|검색/i.test((x.id||'')+' '+(x.className||'')+' '+(x.placeholder||'')+' '+(x.getAttribute('aria-label')||'')))
            )''',10))
            if not ready:
                self.msg("카카오맵 페이지 확인: "+str(info.get("url",""))[:120])
                raise RuntimeError("카카오맵 검색창을 준비하지 못했습니다. Chrome에서 카카오맵 화면이 완전히 열린 뒤 다시 시도해주세요.")
        except Exception as e:
            self.running=False
            self.status.set("시작 실패")
            self.msg("작업 시작 실패: "+str(e))
            return

        for n,row in enumerate(self.rows,1):
            if self.stopflag: break
            self.pb.config(value=n/total*100)
            self.status.set(f'{n}/{total} 처리 중: {row[4]}')
            name=(' ' if self.mode.get()=='space' else '\n').join(row[:4]).strip()
            try:
                res=self.c.js(SEARCH_AND_SELECT_JS.replace('__Q__',js_quote(row[4])),10) or {}
                if not res.get('ok'): raise RuntimeError(res.get('error','주소 검색 실패'))
                self.msg(f'[{n}/{total}] 검색 결과 선택: {res.get("text","")[:100]}')
                fav=self.c.js(CLICK_FAVORITE_JS,5) or {}
                if not fav.get('ok'): raise RuntimeError(fav.get('error','즐겨찾기 버튼을 찾지 못했습니다.'))
                sel=self.c.js(SELECT_GROUP_JS.replace('__GROUP__',js_quote(group_name)),5) or {}
                if sel.get('duplicate'):
                    dup+=1; ow.append(row+['중복','이미 등록된 주소'])
                    self.msg(f'[{n}/{total}] 중복: {row[4]}')
                    self.c.js(CLOSE_LAYERS_JS,5); continue
                if not sel.get('ok'): raise RuntimeError(sel.get('error','그룹 선택 실패'))
                saved=self.c.js(SAVE_FAVORITE_JS.replace('__NAME__',js_quote(name)),5) or {}
                if not saved.get('ok'): raise RuntimeError(saved.get('error','즐겨찾기 저장 실패'))
                ok+=1; ow.append(row+['성공',''])
                self.msg(f'[{n}/{total}] 성공: {row[4]}')
            except Exception as e:
                fail+=1; msg=str(e)
                ew.append(row+[msg]); ow.append(row+['실패',msg])
                self.msg(f'[{n}/{total}] 실패: {row[4]} / {msg}')
                try:self.c.js(CLOSE_LAYERS_JS,5)
                except Exception:pass

        ts=time.strftime('%Y%m%d_%H%M%S')
        rp=RESULT/f'result_{ts}.xlsx'; ep=ERROR/f'error_{ts}.xlsx'
        out.save(rp); er.save(ep)
        self.running=False; self.status.set(f'완료: 성공 {ok} / 중복 {dup} / 실패 {fail}')
        self.msg(f'완료. 결과: {rp}'); self.msg(f'실패목록: {ep}')
        self.r.after(0,lambda:messagebox.showinfo('작업 완료',
            f'성공 {ok}건 / 중복 {dup}건 / 실패 {fail}건\n\n결과: {rp}\n실패목록: {ep}'))

if __name__=='__main__':
    root=tk.Tk(); App(root); root.mainloop()
