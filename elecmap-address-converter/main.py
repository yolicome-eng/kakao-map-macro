import json, re, time, traceback
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox
from openpyxl import load_workbook
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC

URL="https://elecmap.kr/search/single"

def norm(v):
    return re.sub(r"\\s+"," ",str(v or "")).strip()

def looks_like_address(s):
    s=norm(s)
    return bool(re.search(r"(특별시|광역시|특별자치시|특별자치도|도)\\s+.+(시|군|구|읍|면|동|리)",s) or re.search(r"\\b\\d{1,5}-\\d{1,5}\\b",s))

def collect_strings(obj, found):
    if isinstance(obj, dict):
        for k,v in obj.items():
            kl=str(k).lower()
            if any(x in kl for x in ["address","addr","도로명","지번","roadaddress","jibunaddress"]):
                if isinstance(v,str) and looks_like_address(v): found.append(v)
            collect_strings(v,found)
    elif isinstance(obj,list):
        for v in obj: collect_strings(v,found)

def extract_network_address(driver, since):
    for entry in driver.get_log("performance"):
        try:
            msg=json.loads(entry["message"])["message"]
            if msg.get("method")!="Network.responseReceived": continue
            p=msg["params"]; resp=p["response"]; rid=p["requestId"]
            if "elecmap.kr" not in resp.get("url",""): continue
            ct=resp.get("mimeType","")
            if "json" not in ct: continue
            try:
                body=driver.execute_cdp_cmd("Network.getResponseBody",{"requestId":rid})["body"]
                data=json.loads(body); found=[]; collect_strings(data,found)
                if found: return norm(found[0])
            except Exception: pass
        except Exception: pass
    return ""

def find_input(driver):
    els=driver.find_elements(By.TAG_NAME,"input")
    for e in els:
        try:
            ph=(e.get_attribute("placeholder") or "")+" "+(e.get_attribute("aria-label") or "")
            if "전산화번호" in ph or "8자리" in ph: return e
        except Exception: pass
    return els[0] if els else None

def click_search(driver):
    buttons=driver.find_elements(By.TAG_NAME,"button")
    for b in buttons:
        try:
            t=norm(b.text)
            if t=="검색" and b.is_displayed() and b.is_enabled():
                b.click(); return
        except Exception: pass
    raise RuntimeError("검색 버튼을 찾지 못했습니다.")

def run():
    root=tk.Tk(); root.withdraw()
    src=filedialog.askopenfilename(title="전산화번호 엑셀 선택",filetypes=[("Excel","*.xlsx")])
    if not src: return
    out=Path(src).with_name(Path(src).stem+"_주소결과.xlsx")
    wb=load_workbook(src); ws=wb.active
    if ws.max_column<2: ws.cell(1,2,"주소")
    # preserve no-header files: insert address in column B
    for r in range(1,ws.max_row+1):
        if ws.cell(r,2).value is None: ws.cell(r,2,"")
    options=webdriver.ChromeOptions()
    options.add_argument("--disable-gpu"); options.add_argument("--no-sandbox")
    options.set_capability("goog:loggingPrefs",{"performance":"ALL"})
    driver=webdriver.Chrome(options=options)
    driver.set_page_load_timeout(30)
    ok=fail=0
    try:
        driver.get(URL); time.sleep(2)
        for r in range(1,ws.max_row+1):
            code=norm(ws.cell(r,1).value).upper()
            if not re.fullmatch(r"\\d{4}[A-Z]\\d{3}",code): continue
            try:
                inp=find_input(driver)
                if not inp: raise RuntimeError("전산화번호 입력칸을 찾지 못했습니다.")
                inp.click(); inp.send_keys(Keys.CONTROL,"a"); inp.send_keys(code)
                click_search(driver); time.sleep(2)
                addr=extract_network_address(driver,0)
                if not addr:
                    # fallback: visible page text
                    text=driver.find_element(By.TAG_NAME,"body").text
                    candidates=[norm(x) for x in text.splitlines() if looks_like_address(x)]
                    addr=candidates[0] if candidates else ""
                if addr:
                    ws.cell(r,2,addr); ok+=1
                else:
                    ws.cell(r,2,"검색 결과 주소 없음"); fail+=1
            except Exception as e:
                ws.cell(r,2,"오류: "+str(e)[:120]); fail+=1
            wb.save(out)
        messagebox.showinfo("완료",f"완료되었습니다.\\n성공: {ok}건\\n실패: {fail}건\\n\\n{out}")
    except Exception:
        Path("error.log").write_text(traceback.format_exc(),encoding="utf-8")
        messagebox.showerror("오류","오류가 발생했습니다. error.log를 확인하세요.")
    finally:
        driver.quit()

if __name__=="__main__": run()
